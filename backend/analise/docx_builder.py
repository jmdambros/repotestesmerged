"""Converte o rascunho de peça em .docx formatado, com aviso de minuta no topo."""
from __future__ import annotations

import os
import re
import uuid
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

OUTPUT_DIR = Path(os.environ.get("SOPHIA_OUTPUT_DIR", "generated_docs"))
AVISO = "MINUTA GERADA POR IA — REVISAR ANTES DE PROTOCOLAR"
_SAFE_NAME = re.compile(r"^[A-Za-z0-9_\-]+\.docx$")


def build_docx(rascunho: str, tipo_peca: str) -> str:
    """Gera o arquivo e devolve o nome (para o link de download)."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    d.styles["Normal"].font.name = "Times New Roman"
    d.styles["Normal"].font.size = Pt(12)

    aviso = d.add_paragraph()
    aviso.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = aviso.add_run(AVISO)
    run.bold = True

    for line in rascunho.splitlines():
        text = line.strip()
        if not text:
            continue
        clean = re.sub(r"^#+\s*", "", text).replace("**", "")
        is_heading = text.startswith("#") or (clean.isupper() and len(clean) < 90)
        if is_heading:
            d.add_heading(clean, level=2)
        else:
            d.add_paragraph(clean)

    filename = f"{uuid.uuid4().hex}.docx"
    d.save(OUTPUT_DIR / filename)
    return filename


def resolve_download(filename: str) -> Path | None:
    """Valida o nome (sem path traversal) e devolve o caminho, se existir."""
    if not _SAFE_NAME.match(filename):
        return None
    path = OUTPUT_DIR / filename
    return path if path.is_file() else None
